# ---Auto-Refresh v0.2.67 by avalkon--- #
#I know, I have issues with remembering how I was even naming stuff at any given hour.
#sucks, doesn't it? but at least it works. mostly.
import tkinter as tk
import threading
import time
import configparser
import re
from openpyxl import Workbook, load_workbook
from pynput.mouse import Button, Controller
from pynput import keyboard
from python_imagesearch.imagesearch import imagesearch
from playsound3 import playsound
from pathlib import Path

config_file = Path(__file__).parent / "config.ini"
config = configparser.ConfigParser()
config.read(config_file)

# ---Files--- #

def config_path(name):
    path = Path(config["files"][name])
    if path.is_absolute():
        return path
    return config_file.parent / path

nta = config_path("nta")
acqimg = config_path("acquire")
acq2img = config_path("acquire2")
acq3img = config_path("acquire3")
timers = [
    (config_path("submit1"), config.getfloat("confidence", "submit1")),
    (config_path("submit2"), config.getfloat("confidence", "submit2")),
    (config_path("submit3"), config.getfloat("confidence", "submit3")),
    (config_path("submit4"), config.getfloat("confidence", "submit4")),
    (config_path("submit5"), config.getfloat("confidence", "submit5")),]
sound = config_path("sound")
icon = config_path("icon")

# ---Mouse positions--- #
#this is all so you could have a config file. it was all tidier when I hardcoded filepaths.

def get_position(name):
    x, y = config["positions"][name].split(",")
    return int(x), int(y)

refpos = get_position("refresh")
acqpos = get_position("acquire")
acq2pos = get_position("acquire2")
acq3pos = get_position("acquire3")
subpos = get_position("submit")
wigmov = get_position("wiggle")

# ---Hours logging and other miscellanea--- #

refdel1 = config.getfloat("delays", "refresh")
refdel2 = config.getfloat("delays", "refresh2")
acqdel = config.getfloat("delays", "acquire")
waitdel = config.getfloat("delays", "wait")
subdel = config.getfloat("delays", "submit")
wigdel = config.getfloat("delays", "wiggle")

log_folder = Path(__file__).parent / "logs"
log_folder.mkdir(exist_ok=True)
def get_logs():
    today = time.strftime("%Y-%m-%d")
    return log_folder / f"{today}.log"

session_start_time = None
acquire_time = None
last_task_time = None
daily_total = 0
task_count = 0
session_running = False
mouse = Controller()
sound_on = True
auto_submit = False
show_threads = False
pref_door = 1
mouse_move = True
wiggle_time = None
worker_thread = None

stopped = 0
refreshing = 1
acquiring = 2
waiting = 3
submitting = 4
status = stopped

def log_hours(message):
    timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
    log_file = get_logs()
    with open(log_file, "a") as f:
        f.write(f"{timestamp} - {message}\n")

def format_elapsed(seconds):
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    seconds = int(seconds % 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"

def load_daily_stats():
    today = time.strftime("%Y-%m-%d")
    log_file = get_logs()
    daily_total = 0
    task_count = 0
    if log_file.exists():
        with open(log_file, "r") as f:
            for line in f:
                if not line.startswith(today):
                    continue
                if "Task #" not in line:
                    continue
                task_count += 1
                daily_text = (line.split("Today's Tasks': ")[1].split(" | ")[0])
                hours, minutes, seconds = map(int, daily_text.split(":"))
                daily_total = (hours * 3600 + minutes * 60 + seconds)
    return daily_total, task_count
daily_total, task_count = load_daily_stats()

def log_submit():
    global last_task_time
    global daily_total
    global task_count
    global acquire_time
    now = time.monotonic()
    if acquire_time is not None:
        acquire_elapsed = now - acquire_time
        acquire_text = format_elapsed(acquire_elapsed)
    else:
        acquire_text = "--:--:--"
    if last_task_time is not None:
        task_elapsed = now - last_task_time
        daily_total += task_elapsed
        submit_text = format_elapsed(task_elapsed)
    else:
        submit_text = "--:--:--"
    if session_start_time is not None:
        session_elapsed = now - session_start_time
        session_text = format_elapsed(session_elapsed)
    else:
        session_text = "--:--:--"
    task_count += 1
    daily_text = format_elapsed(daily_total)
    log_hours(
        f"Task #{task_count} | "
        f"Since Acquired: {acquire_text} | "
        f"Task Time: {submit_text} | "
        f"Today's Tasks': {daily_text} | "
        f"Session Total: {session_text}")
    last_task_time = now

def update_timers():
    now = time.monotonic()
    if session_running and session_start_time is not None:
        elapsed = now - session_start_time
        session_label.config(text=f"Session: {format_elapsed(elapsed)}")
    if acquire_time is not None:
        elapsed = now - acquire_time
        acquire_timer_label.config(text=f"Since Acquired: {format_elapsed(elapsed)}")
    else:
        acquire_timer_label.config(text="Since Acquired: --:--:--")
    if last_task_time is not None:
        elapsed = now - last_task_time
        task_timer_label.config(text=f"Task Time: {format_elapsed(elapsed)}")
    else:
        task_timer_label.config(text="Task Time: --:--:--")
    daily_total_label.config(text=f"Today's Tasks': {format_elapsed(daily_total)}")
    task_count_label.config(text=f"Tasks: {task_count}")
    root.after(1000, update_timers)

def update_xl():
    week_file = Path(__file__).parent / "totals.xlsx"
    log_folder = Path(__file__).parent / "logs"
    if week_file.exists():
        wb = load_workbook(week_file)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Week"
        ws["A1"] = "Date"
        ws["B1"] = "Day Total"
    for log_file in sorted(log_folder.glob("*.log")):
        match = re.fullmatch(r"(\d{4}-\d{2}-\d{2})\.log", log_file.name)
        if not match:
            continue
        date_text = match.group(1)
        final_total = None
        with open(log_file, "r") as f:
            for line in f:
                match = re.search(r"Today's Tasks':\s*(\d{2}:\d{2}:\d{2})", line)
                if match:
                    final_total = match.group(1)
        if final_total is None:
            continue
        existing_row = None
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=1).value) == date_text:
                existing_row = row
                break
        if existing_row is not None:
            ws.cell(row=existing_row, column=2).value = final_total
        else:
            new_row = ws.max_row + 1
            ws.cell(row=new_row, column=1).value = date_text
            ws.cell(row=new_row, column=2).value = final_total
    wb.save(week_file)
    label(submit_label, "Xl updated")

def set_status(value):
    global status
    status = value

def found(image, confidence=0.8):
    return imagesearch(str(image), confidence)[0] != -1

def submit_found():
    return any(found(image, confidence) for image, confidence in timers)

def sleep(seconds, state):
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        if status != state:
            return False
        time.sleep(0.1)
    return True

def get_acq_pos():
    if found(acq3img):
        if pref_door == 3:
            return acq3pos
        elif pref_door == 2:
            return acq2pos
        else:
            return acqpos
    elif found(acq2img):
        if pref_door == 2:
            return acq2pos
        else:
            return acqpos
    elif found(acqimg):
        return acqpos
    return None

def click(position):
    mouse.position = position
    mouse.click(Button.left)

def label(widget, text):
    root.after(0, lambda: widget.config(text=text))

def play_sound():
    if sound_on:
        playsound(str(sound))

def worker():
    while status != stopped:
        if status == refreshing:
            refresh()
        elif status == acquiring:
            acquire()
        elif status == waiting:
            wait()
        elif status == submitting:
            submit()

def refresh():
    global acquire_time
    global last_task_time
    label(nta_label, "Checking...")
    label(submit_label, "Waiting")
    if not sleep(refdel1, refreshing):
        return
    if found(nta):
        label(nta_label, "NTA")
        if not sleep(refdel2, refreshing):
            return
        click(refpos)
        acquire_time = None
        last_task_time = None
    else:
        label(nta_label, "Not NTA!")
        set_status(acquiring)

def acquire():
    global acquire_time
    global last_task_time
    if not sleep(acqdel, acquiring):
        return
    position = get_acq_pos()
    if position is not None:
        label(nta_label, "Acquiring!")
        label(submit_label, "Waiting")
        click(position)
        acquire_time = time.monotonic()
        last_task_time = time.monotonic()
        log_hours("Acquired")
        play_sound()
        position = None
    else:
        label(nta_label, "Already acquired")
        label(submit_label, "Waiting")
    set_status(waiting)

def wait():
    label(submit_label, "Waiting")
    if found(nta):
        label(nta_label, "NTA")
        set_status(refreshing)
        return
    if found(acqimg):
        label(nta_label, "Acquiring")
        set_status(acquiring)
        return
    if found(acq2img):
        label(nta_label, "Acquiring")
        set_status(acquiring)
        return
    if found(acq3img):
        label(nta_label, "Acquiring")
        set_status(acquiring)
        return
    label(nta_label, "Task")
    sleep(waitdel, waiting)

def submit():
    global wiggle_time
    if wiggle_time is None:
        wiggle_time = time.monotonic()
    if submit_found():
        label(submit_label, "Submitting!")
        log_submit()
        if auto_submit:
            click(subpos)
        if not sleep(acqdel, submitting):
            return
        play_sound()
        set_status(refreshing)
        wiggle_time = None
    else:
        label(submit_label, "Waiting to submit")
        now = time.monotonic()
        if mouse_move and now - wiggle_time >= wigdel:
            mouse.move(*wigmov)
            wiggle_time = now
        if not sleep(subdel, submitting):
            return
        label(submit_label, "Checking...")

def start_worker():
    global worker_thread
    if worker_thread and worker_thread.is_alive():
        return
    worker_thread = threading.Thread(target=worker, daemon=True)
    worker_thread.start()

def start_refresh():
    global session_start_time
    global acquire_time
    global last_task_time
    global session_running
    session_start_time = time.monotonic()
    acquire_time = None
    last_task_time = None
    session_running = True
    set_status(refreshing)
    label(nta_label, "Starting...")
    label(submit_label, "Waiting")
    start_worker()

def start_submit():
    set_status(submitting)
    label(submit_label, "Checking...")
    start_worker()

def stop_all(event=None):
    global session_running
    global acquire_time
    global last_task_time
    session_running = False
    set_status(stopped)
    nta_label.config(text="stopped")
    submit_label.config(text="stopped")
    log_submit()
    acquire_time = None
    last_task_time = None
    if show_threads:
        checkthreads()

def checkthreads():
    threads_label.config(text=f"Threads: {threading.active_count()}")

def open_settings():
    global sound_on
    global auto_submit
    global show_threads
    window = tk.Toplevel(root)
    window.title("Settings")
    window.geometry("260x240")
    window.minsize(260, 240)
    settings_label = tk.Label(window, text="Settings")
    settings_label.grid(row=0, column=0, columnspan=2)

    sound_var = tk.StringVar(value="On" if sound_on else "Off")
    snd_label = tk.Label(window, text="Sound:")
    snd_label.grid(row=1, column=0)
    snd_menu = tk.OptionMenu(window, sound_var, "On", "Off")
    snd_menu.grid(row=1, column=1)

    submit_var = tk.StringVar(value="On" if auto_submit else "Off")
    sub_label = tk.Label( window, text="Auto-Submit Click:")
    sub_label.grid(row=2, column=0)
    sub_menu = tk.OptionMenu(window, submit_var, "On", "Off")
    sub_menu.grid(row=2, column=1)

    threads_var = tk.StringVar(value="On" if show_threads else "Off")
    show_threads_label = tk.Label( window, text="Show Threadcount Button:")
    show_threads_label.grid(row=3, column=0)
    threads_menu = tk.OptionMenu(window, threads_var, "On", "Off")
    threads_menu.grid(row=3, column=1)

    door_var = tk.StringVar(value=str(pref_door))
    door_label = tk.Label( window, text="Preferred Door:")
    door_label.grid(row=4, column=0)
    door_menu = tk.OptionMenu(window, door_var, "1", "2", "3")
    door_menu.grid(row=4, column=1)

    momove_var = tk.StringVar(value="On" if mouse_move else "Off")
    momove_label = tk.Label( window, text="Wiggle while Auto-submit:")
    momove_label.grid(row=5, column=0)
    momove_menu = tk.OptionMenu(window, momove_var, "On", "Off")
    momove_menu.grid(row=5, column=1)

    def save():
        global sound_on
        global auto_submit
        global show_threads
        global pref_door
        global mouse_move
        sound_on = sound_var.get() == "On"
        auto_submit = submit_var.get() == "On"
        show_threads = threads_var.get() == "On"
        pref_door = int(door_var.get())
        mouse_move = momove_var.get() == "On"
        if show_threads:
            thread_button.grid()
            threads_label.grid()
        else:
            thread_button.grid_remove()
            threads_label.grid_remove()
        window.destroy()
    save_button = tk.Button(window, text="Save", width=16, command=save)
    save_button.grid(row=6, column=0, columnspan=2)

root = tk.Tk()
root.title("Auto-Refresh v0.2.67")
root.geometry("337x248+990+540")
root.minsize(337, 248)
root.iconphoto(True, tk.PhotoImage(file=str(icon)))
root.after(1000, update_timers)

# ---Buttons n such--- #

refresh_button = tk.Button(root, text="Start Auto-Refresh", width=16, command=start_refresh)
refresh_button.grid(row=0, column=0)

submit_button = tk.Button(root, text="Auto-Submit", width=16, command=start_submit)
submit_button.grid(row=0, column=1)

nta_label = tk.Label(root, text="stopped", width=16)
nta_label.grid(row=1, column=0)

submit_label = tk.Label(root, text="Waiting", width=16)
submit_label.grid(row=1, column=1)

session_label = tk.Label(root, text="Session: 00:00:00", width=20)
session_label.grid(row=2, column=0)

acquire_timer_label = tk.Label(root, text="Since Acquired: 00:00:00", width=20)
acquire_timer_label.grid(row=2, column=1)

task_timer_label = tk.Label(root, text="Task Time: 00:00:00", width=20)
task_timer_label.grid(row=3, column=0)

task_count_label = tk.Label(root, text=f"Today Submits: {task_count}", width=25)
task_count_label.grid( row=3, column=1, columnspan=2)

daily_total_label = tk.Label(root, text=f"Today's Tasks': {format_elapsed(daily_total)}", width=25)
daily_total_label.grid(row=4, column=0, columnspan=2)

def trigger_acquire():
    global acquire_time
    global last_task_time
    acquire_time = time.monotonic()
    last_task_time = time.monotonic()
    log_hours("Acquired")
    acquire_timer_label.config(text="Since Acquired: 00:00:00")
start_acquire_button = tk.Button(root, text="Trigger Acq Time",width=16,command=trigger_acquire)
start_acquire_button.grid(row=5, column=0)

submit_time_button = tk.Button(root, text="Trigger Sub Time", width=16, command=log_submit)
submit_time_button.grid(row=5, column=1)

stop_button = tk.Button(root, text="Stop All (F8)", width=16, command=stop_all)
stop_button.grid(row=6, column=0)
def stop_all_keybind(event):
    stop_all()
stop_button.bind("<F8>", stop_all_keybind)
def press_stop_key(key):
    if key == keyboard.Key.f8:
        stop_all_keybind(None)
listener = keyboard.Listener(on_press=press_stop_key)
listener.start()
#this should make f8 stop everything no matter what window is active.

settings_button = tk.Button(root, text="Settings", width=16, command=open_settings)
settings_button.grid(row=6, column=1)

export_button = tk.Button(root, text="Update Spreadsheet", width=16, command=update_xl)
export_button.grid(row=7, column=0, columnspan=2)

credit_label = tk.Label(root, text="Auto-Refresh v0.2.67, by avalkon")
credit_label.grid(row=8, column=0, columnspan=2)

thread_button = tk.Button(root, text="Get Threadcount", width=16, command=checkthreads)
thread_button.grid(row=9, column=0)

threads_label = tk.Label(root, text="Threads: --")
threads_label.grid(row=9, column=1)
if not show_threads:
    thread_button.grid_remove()
    threads_label.grid_remove()

def close_program():
    global status
    status = stopped
    root.destroy()
root.protocol("WM_DELETE_WINDOW", close_program)
#it's more of an exit strategy than you'll ever need at this job...
root.mainloop()
